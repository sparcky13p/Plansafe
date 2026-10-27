"""Import a supplier price list (CSV or XLSX) into a Supplier catalogue.

Suppliers use their own column names, so a column mapping (theirs -> ours)
can be passed; the UI mapping wizard (run 6) produces it. Rows that cannot be
mapped are reported, not silently dropped.

Our columns:
  sku_code, name, brand, item, pack_qty, price, specs, in_stock, lead_time_days, kit
  - item:  generic item id (see generic_items.yaml); blank for kits
  - kit:   "item:qty; item:qty" for kits, e.g. "guardrail_post:10; guardrail_top_rail_m:20"
  - specs: "IS 2925; IS 3521"
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

from pydantic import ValidationError

from safetyplan.catalog.model import SKU, Supplier
from safetyplan.rules.engine import generic_items

FIELDS = ("sku_code", "name", "brand", "item", "pack_qty", "price", "specs", "in_stock", "lead_time_days", "kit")
REQUIRED = ("sku_code", "name", "price")


@dataclass
class ImportReport:
    supplier: Supplier
    errors: list[str] = field(default_factory=list)  # "row 7: unknown item 'helmett'"
    rows_read: int = 0


def import_price_list(
    src: str | Path | bytes,
    supplier_id: str,
    supplier_name: str,
    filename: str | None = None,
    columns: dict[str, str] | None = None,
    item_aliases: dict[str, str] | None = None,
) -> ImportReport:
    """columns: their header -> our field. item_aliases: their item text -> generic item id."""
    name = filename or (str(src) if not isinstance(src, bytes) else "upload.csv")
    rows = _read_rows(src, name)
    colmap = {k.strip().lower(): v for k, v in (columns or {}).items()}
    aliases = {k.strip().lower(): v for k, v in (item_aliases or {}).items()}
    known = generic_items()

    skus: list[SKU] = []
    errors: list[str] = []
    for n, row in enumerate(rows, start=2):  # row 1 is the header
        r = {colmap.get(k.strip().lower(), k.strip().lower()): (v or "").strip() for k, v in row.items() if k}
        if not any(r.values()):
            continue
        missing = [f for f in REQUIRED if not r.get(f)]
        if missing:
            errors.append(f"row {n}: missing {', '.join(missing)}")
            continue
        try:
            provides = _provides(r, known, aliases)
            sku = SKU(
                sku_code=r["sku_code"],
                name=r["name"],
                brand=r.get("brand", ""),
                provides=provides,
                price=_num(r["price"]),
                specs=[s.strip() for s in r.get("specs", "").replace(",", ";").split(";") if s.strip()],
                in_stock=r.get("in_stock", "yes").lower() not in ("no", "n", "false", "0", "out of stock"),
                lead_time_days=int(_num(r["lead_time_days"])) if r.get("lead_time_days") else None,
            )
        except (ValueError, ValidationError) as exc:
            errors.append(f"row {n} ({r.get('sku_code')}): {_first_line(exc)}")
            continue
        skus.append(sku)

    codes = [s.sku_code for s in skus]
    dupes = sorted({c for c in codes if codes.count(c) > 1})
    if dupes:
        errors.append(f"duplicate sku_code(s): {', '.join(dupes)}")
    return ImportReport(Supplier(id=supplier_id, name=supplier_name, skus=skus), errors, len(rows))


def _provides(r: dict, known: dict, aliases: dict) -> dict[str, float]:
    if r.get("kit"):
        out: dict[str, float] = {}
        for part in r["kit"].split(";"):
            if not part.strip():
                continue
            item, _, q = part.partition(":")
            item = _item(item, known, aliases)
            out[item] = out.get(item, 0) + _num(q)
        return out
    if not r.get("item"):
        raise ValueError("no item or kit given")
    return {_item(r["item"], known, aliases): _num(r.get("pack_qty") or "1")}


def _item(text: str, known: dict, aliases: dict) -> str:
    t = text.strip()
    if t in known:
        return t
    if t.lower() in aliases:
        return aliases[t.lower()]
    raise ValueError(f"unknown item {t!r}")


def _num(s: str) -> float:
    s = str(s).replace(",", "").replace("₹", "").replace("Rs.", "").replace("Rs", "").strip()
    try:
        return float(s)
    except ValueError:
        raise ValueError(f"not a number: {s!r}") from None


def _read_rows(src, name: str) -> list[dict]:
    data = src if isinstance(src, bytes) else Path(src).read_bytes()
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(it)]
        return [
            {h: ("" if v is None else str(v)) for h, v in zip(header, row)}
            for row in it
        ]
    text = data.decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(text)))


def _first_line(exc: Exception) -> str:
    if isinstance(exc, ValidationError):
        e = exc.errors()[0]
        return f"{'.'.join(map(str, e['loc']))}: {e['msg']}"
    return str(exc).splitlines()[0]
