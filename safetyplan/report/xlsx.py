"""Purchase order workbook: PO sheet for the supplier + requirements and trace."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from safetyplan.rules.engine import generic_items

BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="E6E6E6")
RED = Font(color="B01818")


def build_po(project: dict, supplier_id: str | None = None) -> bytes:
    r = project["result"]
    rec = r["recommendation"]
    sel = supplier_id or project.get("selected_supplier") or rec["best"]
    basket = next((b for b in rec["baskets"] if b["supplier_id"] == sel), None)
    if basket is None:
        raise ValueError(f"no basket for supplier {sel!r}")
    items = generic_items()
    wb = Workbook()

    po = wb.active
    po.title = "Purchase order"
    po.append(["Purchase order"])
    po["A1"].font = Font(bold=True, size=14)
    po.append(["Supplier", basket["supplier_name"]])
    po.append(["Project", project["name"]])
    po.append(["Date", date.today().isoformat()])
    po.append([])
    header = ["SKU code", "Description", "Brand", "Pack", "Qty", "Unit price (INR)", "Line total (INR)"]
    po.append(header)
    hr = po.max_row
    for l in basket["lines"]:
        po.append([l["sku_code"], l["name"], l["brand"], l["pack"], l["packs"], l["unit_price"], None])
        row = po.max_row
        po.cell(row, 7, f"=E{row}*F{row}")
    last = po.max_row
    po.append(["", "", "", "", "", "Total", f"=SUM(G{hr + 1}:G{last})" if last > hr else 0])
    po.cell(po.max_row, 6).font = BOLD
    po.cell(po.max_row, 7).font = BOLD
    for c in range(1, len(header) + 1):
        po.cell(hr, c).font = BOLD
        po.cell(hr, c).fill = HEAD
    for row in po.iter_rows(min_row=hr + 1, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = "#,##0"
    if basket["gaps"]:
        po.append([])
        po.append(["Not available from this supplier"])
        po.cell(po.max_row, 1).font = BOLD
        for g in basket["gaps"]:
            po.append([items[g["item"]]["name"], g["need"], items[g["item"]]["unit"], g["reason"]])
    _widths(po, [14, 60, 14, 10, 8, 16, 16])

    req = wb.create_sheet("Requirements")
    req.append(["Category", "Item", "Unit", "Rule minimum", "Required", "On hand", "To buy",
                "Rule verified", "Below minimum", "Notes"])
    for c in range(1, 11):
        req.cell(1, c).font = BOLD
        req.cell(1, c).fill = HEAD
    for l in r["purchase"]["lines"]:
        req.append([l["category"], l["name"], l["unit"], l["rule_min"], l["required"], l["on_hand"], l["to_buy"],
                    "yes" if l["verified"] else "no", "YES" if l["below_minimum"] else "", "; ".join(l["flags"])])
        if l["below_minimum"]:
            for c in range(1, 11):
                req.cell(req.max_row, c).font = RED
    _widths(req, [18, 44, 6, 12, 10, 9, 9, 12, 13, 70])

    tr = wb.create_sheet("Trace")
    tr.append(["Item", "Feature", "Feature type", "Edges", "Rule", "Pack", "Citation", "Formula", "Inputs",
               "Quantity", "Governing", "Verified"])
    for c in range(1, 13):
        tr.cell(1, c).font = BOLD
        tr.cell(1, c).fill = HEAD
    for rq in r["requirements"]:
        for t in rq["trace"]:
            tr.append([rq["name"], t["feature_id"], t["feature_type"], ",".join(t["edge_ids"]), t["rule_id"],
                       t["pack"], t["citation"], t["formula"],
                       ", ".join(f"{k}={v:g}" for k, v in t["inputs"].items()), t["qty"],
                       "yes" if t["governing"] else "no", "yes" if t["verified"] else "no"])
    _widths(tr, [34, 9, 13, 10, 28, 12, 50, 40, 40, 10, 10, 9])
    for ws in (req, tr):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
