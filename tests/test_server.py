"""Project workflow through the HTTP API on a throwaway SQLite database."""

import io
import json
import time

import pytest

INPUTS = {"floors": 4, "floor_height": 3.0, "workforce": 40}


def new_project(server, **kw):
    r = server.post("/api/projects", json={"name": "Tower A", "inputs": INPUTS, **kw})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def upload(server, pid, fixtures, name="tower_mm.dxf"):
    r = server.post(f"/api/projects/{pid}/plan-file",
                    files={"file": (name, fixtures[name].read_bytes(), "application/octet-stream")})
    assert r.status_code == 200, r.text
    return r.json()


def test_me_reports_dev_mode_and_fallbacks(server):
    body = server.get("/api/me").json()
    assert body["user"] == "dev-user" and body["auth_mode"] == "dev"
    assert body["services"] == {"database": "sqlite", "blob_uploads": False,
                                "vision_model": "mock", "photo_check": False}


def test_full_workflow(server, fixtures):
    pid = new_project(server)
    assert [p["id"] for p in server.get("/api/projects").json()] == [pid]

    proj = upload(server, pid, fixtures)
    assert proj["geometry"]["outline_perimeter"] == pytest.approx(72.4)
    assert proj["geometry"]["view"]["width"] > 100
    assert proj["detection"]["detector"] == "mock"
    bg = server.get(f"/api/projects/{pid}/background.png")
    assert bg.status_code == 200 and bg.content.startswith(b"\x89PNG")

    cut = next(c["id"] for c in proj["geometry"]["candidates"] if "OPENING_CUTOUT" in c["blocks"])
    review = {"types": {cut: "other"}, "protected_edges": ["E7"],
              "added": [{"id": "U1", "type": "floor_opening", "polygon": [[1, 1], [1.8, 1], [1.8, 1.8], [1, 1.8]]}]}
    r = server.put(f"/api/projects/{pid}/review", json=review)
    assert r.status_code == 200, r.text

    parsed = server.post(f"/api/projects/{pid}/instructions/parse",
                         json={"text": "We already have 25 helmets\nsite near school"}).json()
    assert not parsed["instructions_confirmed"]
    r = server.put(f"/api/projects/{pid}/instructions", json=parsed["instructions"])
    assert r.json()["instructions_confirmed"]

    res = server.post(f"/api/projects/{pid}/compute")
    assert res.status_code == 200, res.text
    result = res.json()["result"]
    lines = {l["item"]: l for l in result["purchase"]["lines"]}
    assert lines["guardrail_top_rail_m"]["rule_min"] == 151  # E7 (north, 20 m) protected
    assert lines["hole_cover_1_5m"]["rule_min"] == 4  # cut-out dropped, drawn U1 0.8 m -> 1.1 m cover
    assert lines["helmet"]["to_buy"] == 19
    assert result["instructions_confirmed"] and res.json()["selected_supplier"] == "demo-a"

    other = next(b["supplier_id"] for b in result["recommendation"]["baskets"] if b["supplier_id"] != "demo-a")
    assert server.put(f"/api/projects/{pid}/supplier", json={"supplier_id": other}).json()["selected_supplier"] == other

    pdf = server.get(f"/api/projects/{pid}/report.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    import pymupdf

    text = "".join(p.get_text() for p in pymupdf.open(stream=pdf.content, filetype="pdf"))
    assert "Equipment required" in text and "Sign-off" in text and "OSHA-EDGE-GUARDRAIL" in text

    xl = server.get(f"/api/projects/{pid}/po.xlsx")
    assert xl.status_code == 200
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xl.content))
    po = wb["Purchase order"]
    assert po["B2"].value.startswith("SiteGuard")
    assert any(isinstance(c.value, str) and c.value.startswith("=SUM(") for c in po["G"])
    assert {"Requirements", "Trace"} <= set(wb.sheetnames)

    actions = [e["action"] for e in server.get(f"/api/projects/{pid}/audit").json()]
    assert actions == ["created", "analyzed", "reviewed", "instructions_confirmed", "computed", "supplier_selected"]


def test_compute_requires_confirmed_instructions_when_text_given(server, fixtures):
    pid = new_project(server, instructions_text="harness nahi chahiye")
    upload(server, pid, fixtures)
    r = server.post(f"/api/projects/{pid}/compute")
    assert r.status_code == 409 and "Confirm" in r.json()["detail"]


def test_new_upload_resets_review_and_result(server, fixtures):
    pid = new_project(server)
    upload(server, pid, fixtures)
    server.put(f"/api/projects/{pid}/review", json={"protected_edges": ["E1"]})
    server.post(f"/api/projects/{pid}/compute")
    proj = upload(server, pid, fixtures, "flat_m.dxf")
    assert proj["review"] == {} and proj["result"] is None


def test_review_validation(server, fixtures):
    pid = new_project(server)
    upload(server, pid, fixtures)
    assert server.put(f"/api/projects/{pid}/review", json={"types": {"C99": "stair"}}).status_code == 422
    assert server.put(f"/api/projects/{pid}/review", json={"types": {"C1": "ladder"}}).status_code == 422
    bad_added = {"added": [{"id": "U1", "type": "balcony", "polygon": [[0, 0], [1, 0], [1, 1]]}]}
    assert server.put(f"/api/projects/{pid}/review", json=bad_added).status_code == 422


def test_projects_are_private(server, fixtures):
    from safetyplan.db.models import Project, get_session

    with get_session() as db:
        p = Project(owner="someone-else", name="Theirs", inputs=INPUTS)
        db.add(p)
        db.commit()
        pid = p.id
    assert server.get(f"/api/projects/{pid}").status_code == 404
    assert pid not in [p["id"] for p in server.get("/api/projects").json()]


def test_pdf_plan_with_detected_scale_then_calibration(server, fixtures):
    pid = new_project(server)
    proj = upload(server, pid, fixtures, "tower_pdf.pdf")
    assert proj["plan_file"]["pdf"]["pages"][0]["detected_scale"] == 100
    assert proj["geometry"]["outline_perimeter"] == pytest.approx(70.0, abs=0.05)
    # Calibrate: the 20 m slab edge measured on paper (566.93 pt) is really 10 m -> everything halves.
    k = 72 / 25.4 * 10
    cal = {"x1": 100, "y1": 700, "x2": 100 + 20 * k, "y2": 700, "distance_m": 10}
    proj = server.post(f"/api/projects/{pid}/scale", json={"calibration": cal}).json()
    assert proj["geometry"]["outline_perimeter"] == pytest.approx(35.0, abs=0.05)
    proj = server.post(f"/api/projects/{pid}/scale", json={"scale_ratio": 100, "calibration": None}).json()
    assert proj["geometry"]["outline_perimeter"] == pytest.approx(70.0, abs=0.05)
    # Two-point factor: "this line is really twice as long" -> back to 140 m.
    proj = server.post(f"/api/projects/{pid}/scale", json={"factor": 2.0}).json()
    assert proj["geometry"]["outline_perimeter"] == pytest.approx(140.0, abs=0.1)


def test_dxf_rescale_factor(server, fixtures):
    pid = new_project(server)
    upload(server, pid, fixtures, "flat_m.dxf")  # unitless, assumed metres
    proj = server.post(f"/api/projects/{pid}/scale", json={"factor": 0.5}).json()
    assert proj["geometry"]["outline_perimeter"] == pytest.approx(22.0, abs=0.01)
    assert proj["plan_file"]["unit_scale"] == 0.5


def test_plan_file_url_must_be_vercel_blob(server):
    pid = new_project(server)
    r = server.post(f"/api/projects/{pid}/plan-file-url", json={"url": "http://169.254.169.254/x.dxf", "name": "x.dxf"})
    assert r.status_code == 422


def test_rejects_non_plan_upload(server):
    pid = new_project(server)
    r = server.post(f"/api/projects/{pid}/plan-file", files={"file": ("plan.png", b"x", "image/png")})
    assert r.status_code == 415


def test_invalid_inputs_rejected(server):
    r = server.post("/api/projects", json={"name": "x", "inputs": {"floors": 0, "floor_height": 3, "workforce": 1}})
    assert r.status_code == 422


class TestSuppliers:
    CSV = ("Item Code,Description,Make,Category,Pack Qty,Rate,IS Code,Stock\n"
           "Z-H10,Helmet x10,Zed,Safety Helmet,10,1500,IS 2925,yes\n"
           "Z-W,Widget,Zed,Widget,1,10,,yes\n").encode()

    def test_preview_suggests_mapping(self, server):
        r = server.post("/api/suppliers/preview", files={"file": ("zed.csv", self.CSV, "text/csv")})
        body = r.json()
        assert body["suggested_columns"] == {
            "Item Code": "sku_code", "Description": "name", "Make": "brand", "Category": "item",
            "Pack Qty": "pack_qty", "Rate": "price", "IS Code": "specs", "Stock": "in_stock"}
        assert body["suggested_items"] == {"Safety Helmet": "helmet", "Widget": None}

    def test_saved_supplier_is_used_in_plans(self, server, fixtures):
        preview = server.post("/api/suppliers/preview", files={"file": ("zed.csv", self.CSV, "text/csv")}).json()
        aliases = {k: v for k, v in preview["suggested_items"].items() if v}
        r = server.post("/api/suppliers", files={"file": ("zed.csv", self.CSV, "text/csv")},
                        data={"supplier_name": "Zed Traders", "columns": json.dumps(preview["suggested_columns"]),
                              "item_aliases": json.dumps(aliases)})
        assert r.status_code == 201 and r.json()["skus"] == 1 and "Widget" in r.json()["errors"][0]
        sid = r.json()["id"]
        listed = server.get("/api/suppliers").json()
        assert listed[0]["id"] == sid and listed[0]["own"] and len(listed) == 3

        pid = new_project(server)
        upload(server, pid, fixtures)
        result = server.post(f"/api/projects/{pid}/compute").json()["result"]
        zed = next(b for b in result["recommendation"]["baskets"] if b["supplier_id"] == sid)
        assert [l["sku_code"] for l in zed["lines"]] == ["Z-H10"]

        assert server.delete(f"/api/suppliers/{sid}").status_code == 204
        assert len(server.get("/api/suppliers").json()) == 2


class TestPhotoCheck:
    def _computed(self, server, fixtures):
        pid = new_project(server)
        upload(server, pid, fixtures)
        server.post(f"/api/projects/{pid}/compute")
        return pid

    def test_not_configured(self, server, fixtures):
        pid = self._computed(server, fixtures)
        r = server.post(f"/api/projects/{pid}/photos", files={"file": ("p.jpg", b"jpg", "image/jpeg")},
                        data={"feature_ids": '["E7"]'})
        assert r.status_code == 503

    def test_findings_with_mock_sam(self, server, fixtures, monkeypatch):
        from safetyplan.photo import sam
        from safetyplan.server import app as app_mod

        masks = [
            sam.Mask("slab edge", 0.9, (0.1, 0.5, 0.9, 0.6)),
            sam.Mask("guardrail", 0.8, (0.1, 0.4, 0.5, 0.55)),
            sam.Mask("lift shaft opening", 0.85, (0.6, 0.1, 0.8, 0.3)),
        ]
        monkeypatch.setattr(app_mod, "sam_from_env", lambda: sam.MockSam(masks))
        pid = self._computed(server, fixtures)
        lift = next(f["id"] for f in server.get(f"/api/projects/{pid}").json()["result"]["features"]
                    if f["type"] == "lift_shaft")
        r = server.post(f"/api/projects/{pid}/photos", files={"file": ("p.jpg", b"jpg", "image/jpeg")},
                        data={"feature_ids": json.dumps(["E7", lift, "C2"]), "floor": "2"})
        assert r.status_code == 200, r.text
        status = {f["feature_id"]: f["status"] for f in r.json()["findings"]}
        assert status == {"E7": "protection_visible", lift: "possibly_unprotected", "C2": "not_visible"}
        assert len(server.get(f"/api/projects/{pid}").json()["photo_checks"]) == 1


class TestClerkAuth:
    @pytest.fixture
    def keys(self, monkeypatch):
        import jwt
        from cryptography.hazmat.primitives.asymmetric import rsa

        from safetyplan.config import settings
        from safetyplan.server import auth

        priv = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        monkeypatch.setenv("CLERK_ISSUER", "https://test.clerk.accounts.dev")
        monkeypatch.setenv("CLERK_AUTHORIZED_PARTIES", "https://plansafe.app")
        settings.cache_clear()

        class FakeJwks:
            def get_signing_key_from_jwt(self, token):
                return type("K", (), {"key": priv.public_key()})()

        monkeypatch.setattr(auth, "_jwks", lambda issuer: FakeJwks())

        def token(**claims):
            base = {"sub": "user_123", "iss": "https://test.clerk.accounts.dev", "exp": int(time.time()) + 60,
                    "azp": "https://plansafe.app"}
            return jwt.encode({**base, **claims}, priv, algorithm="RS256")

        yield token
        settings.cache_clear()

    def test_valid_token(self, server, keys):
        r = server.get("/api/me", headers={"Authorization": f"Bearer {keys()}"})
        assert r.status_code == 200 and r.json()["user"] == "user_123" and r.json()["auth_mode"] == "clerk"

    @pytest.mark.parametrize("claims", [{"exp": int(time.time()) - 10}, {"iss": "https://evil"},
                                        {"azp": "https://evil.app"}])
    def test_rejected_tokens(self, server, keys, claims):
        assert server.get("/api/me", headers={"Authorization": f"Bearer {keys(**claims)}"}).status_code == 401

    def test_missing_token(self, server, keys):
        assert server.get("/api/projects").status_code == 401


def test_deployed_without_auth_refuses(server, monkeypatch):
    from safetyplan.config import settings

    monkeypatch.setenv("VERCEL", "1")
    settings.cache_clear()
    assert server.get("/api/projects").status_code == 503
