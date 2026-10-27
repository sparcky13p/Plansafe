import pytest

from safetyplan.catalog.demo import demo_suppliers
from safetyplan.catalog.importer import import_price_list
from safetyplan.catalog.model import SKU


def test_demo_catalogues_import_cleanly():
    a, b = demo_suppliers()
    assert len(a.skus) == 24 and len(b.skus) == 13
    kit = next(s for s in a.skus if s.sku_code == "A-GR-KIT10")
    assert kit.is_kit and kit.provides["guardrail_post"] == 10


def test_supplier_columns_and_item_aliases_mapped():
    _, b = demo_suppliers()
    hel = next(s for s in b.skus if s.sku_code == "B-HEL20")
    assert hel.provides == {"helmet": 20} and hel.price == 4100 and hel.specs == ["IS 2925"]
    assert not next(s for s in b.skus if s.sku_code == "B-COV15").in_stock


def test_bad_rows_are_reported_not_dropped_silently():
    csv = (
        "sku_code,name,item,pack_qty,price\n"
        "X1,Helmet,helmet,10,2000\n"
        "X2,Helmett,helmett,1,200\n"  # unknown item
        "X3,No price,helmet,1,\n"  # missing price
        "X4,Bad price,helmet,1,abc\n"
        "X1,Dupe,helmet,5,900\n"
        ",,,,\n"  # blank row ignored
    ).encode()
    rep = import_price_list(csv, "x", "X", filename="x.csv")
    assert [s.sku_code for s in rep.supplier.skus] == ["X1", "X1"]
    errs = " | ".join(rep.errors)
    assert "row 3" in errs and "unknown item 'helmett'" in errs
    assert "row 4: missing price" in errs
    assert "row 5" in errs and "not a number" in errs
    assert "duplicate sku_code(s): X1" in errs


def test_price_formats():
    csv = "sku_code,name,item,price\nP1,Helmet,helmet,\"Rs. 1,250\"\nP2,Helmet,helmet,₹990\n".encode()
    rep = import_price_list(csv, "p", "P", filename="p.csv")
    assert [s.price for s in rep.supplier.skus] == [1250, 990]


def test_xlsx_import(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["sku_code", "name", "item", "pack_qty", "price", "specs"])
    ws.append(["H10", "Helmet x10", "helmet", 10, 2100, "IS 2925"])
    ws.append(["K1", "Kit", None, None, 5000, None])
    ws["G1"] = "kit"
    ws["G3"] = "helmet:5; reflective_vest:5"
    path = tmp_path / "list.xlsx"
    wb.save(path)
    rep = import_price_list(path, "x", "X")
    assert not rep.errors, rep.errors
    assert rep.supplier.skus[1].provides == {"helmet": 5, "reflective_vest": 5}


@pytest.mark.parametrize("spec,ok", [("IS 3521", True), ("IS-3521:1999", True), ("is3521", True), ("IS 2925", False)])
def test_spec_matching_normalises(spec, ok):
    s = SKU(sku_code="H", name="h", provides={"full_body_harness": 1}, price=1, specs=[spec])
    assert s.meets(["IS 3521"]) is ok
